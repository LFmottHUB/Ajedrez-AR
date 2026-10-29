from pathlib import Path
import fitz
from PIL import Image
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent
(root/'output').mkdir(exist_ok=True)
pdf=fitz.open(root/'printing/selector.pdf')
pix=pdf[0].get_pixmap(matrix=fitz.Matrix(16,16),colorspace=fitz.csGRAY,alpha=False)
im=Image.frombytes('L',(pix.width,pix.height),pix.samples)
# Normalize the marker's outer square, excluding the PDF's thin white page edge.
mask=im.point(lambda v:255 if v<128 else 0)
box=mask.getbbox()
im=im.crop(box).point(lambda v:0 if v<128 else 255, mode='1')
im.save(root/'output/selector-print.png',dpi=(600,600))
d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Cm(1.6)
sec.left_margin=sec.right_margin=Cm(1.8)
for name in ['Normal','Title','Heading 1','Heading 2','Subtitle']:
 s=d.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
for style in d.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.styles['Normal'].font.size=Pt(10)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Title'].font.size=Pt(23)
d.add_paragraph('Marcadores para AR Chess Android','Title')
d.add_paragraph('Guía de impresión y pruebas con un único patrón', 'Subtitle')
d.add_paragraph('Este documento contiene el mismo marcador selector en tres tamaños: 5, 10 y 15 cm. Imprime primero el de 10 cm y úsalo con la APK de prueba. Los tamaños son opciones para evaluar el reconocimiento; no garantizan detección a cualquier distancia o ángulo.')
d.add_heading('Configuración de impresión',1)
for s in [
 'Papel Carta de 21,59 × 27,94 cm, orientación vertical. En Word y en la impresora selecciona el mismo tamaño de papel.',
 'Escala 100 % o Tamaño real. Una página por hoja. Desactiva Ajustar a página, Reducir y cualquier ampliación automática.',
 'Negro puro sobre blanco. El patrón está preparado sin colores ni filtros; no necesita aumentar la saturación. Usa blanco y negro y calidad normal o alta, preferiblemente 600 ppp. Desactiva modo borrador y ahorro de tóner.',
 'Usa papel blanco mate, plano y opaco. Evita papel brillante, plastificado y reflejos. Si lo montas sobre cartón, evita curvarlo.',
 'Conserva al menos 10 mm de margen blanco alrededor del borde negro. No recortes por el borde negro ni tapes sus esquinas.',
 'Después de imprimir, mide el lado exterior negro con una regla. Debe medir exactamente lo indicado; el margen blanco no forma parte de esa medida.'
]:d.add_paragraph(s,style='List Bullet')
d.add_heading('Cómo realizar la prueba',1)
d.add_paragraph('Abre la APK, permite el acceso a la cámara y enfoca un solo marcador completo. Comienza con luz uniforme, el teléfono estable y un ángulo cercano al perpendicular. Acerca o aleja lentamente hasta que quede enfocado. Prueba después otros ángulos y distancias, sin tapar el patrón con los dedos.')
d.add_paragraph('Anota modelo de teléfono, versión Android, tamaño impreso, distancia aproximada, iluminación y si el tablero tiembla o desaparece. Si pierdes el marcador, vuelve a enfocarlo: la partida debe conservarse. No existe una distancia mínima o máxima validada todavía.')
d.add_paragraph('El tablero usa una escala relativa al marcador. No es una herramienta de medición física. Imprime las páginas siguientes por separado; no necesitas los tres marcadores al mismo tiempo.')
d.add_paragraph('Origen del patrón: carlosmccosta/AR-Chess, licencia MIT. Se conserva el dibujo del PDF selector aportado por el usuario. Proyecto: https://github.com/carlosmccosta/AR-Chess',style='Normal')
for cm in [5,10,15]:
 d.add_page_break()
 d.add_heading(f'Marcador de {cm} cm',0)
 d.add_paragraph(f'Lado exterior negro de {cm} × {cm} cm · Imprimir al 100 % · Papel Carta')
 p=d.add_paragraph();p.paragraph_format.space_before=Pt(30);p.paragraph_format.space_after=Pt(30);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pic=p.add_run().add_picture(str(root/'output/selector-print.png'),width=Cm(cm),height=Cm(cm))
 pic._inline.docPr.set('descr',f'Marcador selector negro y blanco, cuadrado exterior de {cm} centímetros por lado.')
 d.add_paragraph(f'Comprobación con regla: el borde exterior negro debe medir {cm} cm en horizontal y {cm} cm en vertical. Conserva al menos 1 cm de blanco libre alrededor. No deformes ni recortes el dibujo.')
settings=d.settings.element
el=OxmlElement('w:doNotAutoCompressPictures');settings.append(el)
d.core_properties.title='Marcadores para imprimir de AR Chess Android'
d.core_properties.subject='Patrón selector a 5 10 y 15 centímetros'
d.save(root/'output/Marcadores_AR_Chess.docx')
print('Created Word with marker raster', im.size, 'crop',box)
